"""Helpers for academic Vietnamese Word reports (Times New Roman, A4, 1.5 spacing)."""

from __future__ import annotations

import hashlib
import io
import re
from pathlib import Path

import requests
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GREEN = RGBColor(0x1B, 0x5E, 0x20)
GREEN2 = RGBColor(0x2E, 0x7D, 0x32)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x33, 0x33, 0x33)
CACHE_DIR = Path(__file__).resolve().parent / "_diagram_cache"
CACHE_DIR.mkdir(exist_ok=True)


def set_run_font(run, name="Times New Roman", size=13, bold=False, italic=False, color=BLACK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_paragraph_format(p, *, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=None, space_before=0, space_after=6, line=1.5):
    pf = p.paragraph_format
    pf.alignment = alignment
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)
    else:
        pf.first_line_indent = Cm(0)


def shade_cell(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "9E9E9E")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Trang ")
    set_run_font(run, size=11, color=GRAY)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r2 = p.add_run()
    r2._r.append(fld1)
    r2._r.append(instr)
    r2._r.append(fld2)
    set_run_font(r2, size=11, color=GRAY)


def init_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    set_page_number(section)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return doc


def add_title(doc, text, size=18, space_after=6):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=space_after, line=1.15)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, color=GREEN)
    return p


def add_center(doc, text, size=13, bold=False, italic=False, space_after=4):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=space_after, line=1.15)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_h(doc, text, level=1):
    sizes = {1: 16, 2: 14, 3: 13, 4: 13}
    colors = {1: GREEN, 2: GREEN2, 3: BLACK, 4: BLACK}
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=14 if level == 1 else 10, space_after=8, line=1.15)
    run = p.add_run(text)
    set_run_font(run, size=sizes.get(level, 13), bold=True, color=colors.get(level, BLACK))
    return p


def add_p(doc, text, *, indent=True):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=1.27 if indent else 0, space_after=8)
    _add_mixed_runs(p, text, size=13)
    return p


def _add_mixed_runs(p, text, size=13):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=size - 1)
        else:
            run = p.add_run(part)
            set_run_font(run, size=size)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=0, space_after=3, line=1.3)
    p.paragraph_format.left_indent = Cm(1.0 + level * 0.75)
    run = p.add_run("• ")
    set_run_font(run, size=13, color=GREEN2)
    _add_mixed_runs(p, text, size=13)
    return p


def add_caption(doc, text, kind="Hình"):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=12, line=1.15)
    run = p.add_run(text)
    set_run_font(run, size=12, italic=True, color=GRAY)
    return p


def add_code(doc, text, lang=""):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_line=0, space_before=4, space_after=8, line=1.15)
    p.paragraph_format.left_indent = Cm(0.4)
    run = p.add_run(text)
    set_run_font(run, name="Consolas", size=9, color=GRAY)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F5F5F5")
    shading.set(qn("w:val"), "clear")
    p._p.get_or_add_pPr().append(shading)
    return p


def add_table(doc, headers, rows, caption=None):
    if caption:
        add_caption(doc, caption, "Bảng")
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, space_before=3, space_after=3, line=1.1)
        run = p.add_run(str(h))
        set_run_font(run, size=10, bold=True, color=WHITE)
        shade_cell(hdr[i], "2E7D32")
        set_cell_border(hdr[i])
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        bg = "E8F5E9" if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            p = cells[c_i].paragraphs[0]
            set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_line=0, space_before=2, space_after=2, line=1.1)
            run = p.add_run(str(val))
            set_run_font(run, size=10)
            shade_cell(cells[c_i], bg)
            set_cell_border(cells[c_i])
    p = doc.add_paragraph()
    set_paragraph_format(p, space_after=8, line=1.0)
    return table


def _image_ext(content: bytes) -> str | None:
    if content[:8] == b"\x89PNG\r\n\x1a\n":
        return ".png"
    if content[:3] == b"\xff\xd8\xff":
        return ".jpg"
    return None


def render_mermaid(source: str) -> Path | None:
    import base64

    key = hashlib.sha1(source.encode("utf-8")).hexdigest()[:16]
    for ext in (".png", ".jpg"):
        cached = CACHE_DIR / f"{key}{ext}"
        if cached.exists() and cached.stat().st_size > 100:
            print(f"  cache hit {cached.name}", flush=True)
            return cached

    def save(content: bytes) -> Path | None:
        ext = _image_ext(content)
        if not ext:
            return None
        out = CACHE_DIR / f"{key}{ext}"
        out.write_bytes(content)
        return out

    try:
        print(f"  mermaid.ink {key}...", flush=True)
        b64 = base64.urlsafe_b64encode(source.encode("utf-8")).decode("ascii")
        url = f"https://mermaid.ink/img/{b64}"
        r = requests.get(url, timeout=25)
        if r.status_code == 200:
            path = save(r.content)
            if path:
                print(f"  mermaid.ink ok {path.name} {len(r.content)}", flush=True)
                return path
        print(f"  mermaid.ink fail {r.status_code} {r.headers.get('content-type')}", flush=True)
    except Exception as exc:
        print(f"  mermaid.ink error {exc}", flush=True)

    try:
        print(f"  kroki {key}...", flush=True)
        r = requests.post(
            "https://kroki.io/mermaid/png",
            data=source.encode("utf-8"),
            headers={"Content-Type": "text/plain"},
            timeout=25,
        )
        if r.status_code == 200:
            path = save(r.content)
            if path:
                print(f"  kroki ok {path.name} {len(r.content)}", flush=True)
                return path
        print(f"  kroki fail {r.status_code}", flush=True)
    except Exception as exc:
        print(f"  kroki error {exc}", flush=True)
    return None


def add_figure(doc, mermaid_src: str, caption: str):
    print("FIGURE", caption.encode("ascii", "replace").decode("ascii"), flush=True)
    path = render_mermaid(mermaid_src)
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, space_before=8, space_after=4, line=1.0)
    if path:
        run = p.add_run()
        try:
            run.add_picture(str(path), width=Cm(16.0))
        except Exception:
            add_code(doc, mermaid_src)
    else:
        add_code(doc, mermaid_src)
    add_caption(doc, caption)
    return p
